"""Journal d'audit : atteignable, filtrable, et réservé à l'administrateur.

La page existait déjà mais restait introuvable — aucun lien n'y menait — et
ses filtres s'affichaient sans jamais être appliqués. Ces tests verrouillent
les trois points : on y accède, les filtres filtrent, et seuls les
administrateurs y entrent.
"""
import io
from datetime import date, timedelta

import app as application
import openpyxl

# Jeu d'essai : (nom consigné, action, entité, détails, décalage en jours).
ENTREES = (
    ('admin',     'LOGIN',           None,       'AUDIT login admin',    0),
    ('admin',     'CREATE_USER',     'user',     'AUDIT création',       0),
    ('rh',        'UPDATE_EMPLOYE',  'employe',  'AUDIT modification',   1),
    ('rh',        'DELETE_DOCUMENT', 'document', 'AUDIT suppression',    2),
    ('manager',   'CLOCK_IN',        'presence', 'AUDIT pointage',       3),
    # Nom consigné sans compte correspondant : il doit rester affiché.
    ('disparu',   'LOGIN',           None,       'AUDIT compte supprimé', 4),
)


def _semer():
    """Insère le jeu d'essai. La table est purgée avant chaque test."""
    aujourd_hui = date.today()
    with application.db_cursor(commit=True) as (conn, cur):
        for username, action, entite, details, decalage in ENTREES:
            cur.execute(
                """INSERT INTO audit_logs
                       (username, action, entity_type, entity_id, details,
                        ip_address, timestamp)
                   VALUES (%s, %s, %s, 1, %s, '127.0.0.1',
                           %s + INTERVAL '12 hours')""",
                (username, action, entite, details,
                 aujourd_hui - timedelta(days=decalage)))


def test_journal_est_reserve_a_ladministrateur(app):
    """Seul l'administrateur entre dans le journal."""
    _semer()
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    assert admin.get('/audit').status_code == 200

    for username, mot_de_passe in (('rh', 'rh123'),
                                   ('manager', 'manager123'),
                                   ('employe', 'user123')):
        client = app.test_client()
        client.post('/login', data={'username': username, 'password': mot_de_passe})
        assert client.get('/audit', follow_redirects=False).status_code in (301, 302), \
            username


def test_lien_du_journal_est_invisible_hors_admin(app):
    """Le lien n'apparaît ni dans le menu ni sur la page des comptes.

    Attention : le rôle RH contourne l'ACL (accès global), le masquage
    automatique des liens ne peut donc pas le masquer. Le gabarit doit le
    conditionner explicitement au rôle administrateur.
    """
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    page_admin = admin.get('/utilisateurs').data.decode('utf-8')
    assert '/audit' in page_admin

    rh = app.test_client()
    rh.post('/login', data={'username': 'rh', 'password': 'rh123'})
    page_rh = rh.get('/utilisateurs').data.decode('utf-8')
    assert '/audit' not in page_rh


def test_chaque_compte_ouvre_son_historique(app):
    """Depuis les comptes, l'historique pointe sur le journal filtré."""
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    page = admin.get('/utilisateurs').data.decode('utf-8')
    assert 'audit?username=' in page

    _semer()
    page = admin.get('/audit?username=rh').data.decode('utf-8')
    assert 'AUDIT modification' in page
    assert 'AUDIT création' not in page


def test_les_filtres_sont_reellement_appliques(app):
    """Les filtres filtrent — ils se contentaient autrement de s'afficher."""
    _semer()
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})

    par_action = admin.get('/audit?action=LOGIN').data.decode('utf-8')
    assert 'AUDIT login admin' in par_action
    assert 'AUDIT création' not in par_action

    par_utilisateur = admin.get('/audit?username=rh').data.decode('utf-8')
    assert 'AUDIT modification' in par_utilisateur
    assert 'AUDIT création' not in par_utilisateur

    par_entite = admin.get('/audit?entity_type=employe').data.decode('utf-8')
    assert 'AUDIT modification' in par_entite
    assert 'AUDIT suppression' not in par_entite


def test_filtre_sur_la_periode(app):
    """« Depuis » écarte les jours plus anciens."""
    _semer()
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    depuis = (date.today() - timedelta(days=1)).isoformat()

    page = admin.get(f'/audit?date_from={depuis}').data.decode('utf-8')
    assert 'AUDIT modification' in page        # hier
    assert 'AUDIT suppression' not in page     # avant-hier


def test_le_nom_consigne_survit_a_la_suppression_du_compte(app):
    """Le nom du log prime sur celui du compte, qui a pu disparaître."""
    _semer()
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})

    page = admin.get('/audit?username=disparu').data.decode('utf-8')
    assert 'disparu' in page
    assert 'AUDIT compte supprimé' in page
    # Et il n'est pas remplacé par « Système ».
    assert 'Système' not in page


def test_le_journal_est_pagine(app):
    """On ne tronque plus à 150 lignes : la liste est paginée."""
    with application.db_cursor(commit=True) as (conn, cur):
        for i in range(60):
            cur.execute(
                """INSERT INTO audit_logs (username, action, details, timestamp)
                   VALUES ('admin', 'LOGIN', %s, CURRENT_TIMESTAMP)""",
                (f'AUDIT masse {i}',))

    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    page = admin.get('/audit').data.decode('utf-8')
    assert page.count('AUDIT masse') <= 25
    assert admin.get('/audit?page=3').status_code == 200


def test_le_journal_ne_propose_aucune_purge(app):
    """Le bouton « Nettoyer » était affiché sans être câblé : il est retiré.

    Purger une piste d'audit est de toute façon discutable — et le faire
    croire sans que rien ne se passe l'est plus encore.
    """
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    assert 'Nettoyer' not in admin.get('/audit').data.decode('utf-8')


def test_export_excel_du_journal(app):
    """L'export produit un classeur exploitable."""
    _semer()
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    response = admin.get('/audit/export/excel')

    assert response.status_code == 200
    assert 'spreadsheetml' in response.headers['Content-Type']
    assert 'attachment' in response.headers['Content-Disposition']
    assert 'journal_audit_' in response.headers['Content-Disposition']

    feuille = openpyxl.load_workbook(io.BytesIO(response.data)).active
    valeurs = ' '.join(str(c.value or '') for ligne in feuille.iter_rows()
                       for c in ligne)
    for marqueur in ('AUDIT login admin', 'AUDIT création', 'AUDIT modification',
                     'AUDIT suppression', 'AUDIT pointage'):
        assert marqueur in valeurs
    # Le fichier rappelle sa sélection : sinon on ne saurait pas à quoi il
    # correspond.
    assert 'Filtres :' in valeurs


def test_export_reprend_les_filtres_de_la_page(app):
    """Un export qui ignorerait la sélection serait trompeur."""
    _semer()
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})

    response = admin.get('/audit/export/excel?username=rh')
    feuille = openpyxl.load_workbook(io.BytesIO(response.data)).active
    valeurs = ' '.join(str(c.value or '') for ligne in feuille.iter_rows()
                       for c in ligne)
    assert 'AUDIT modification' in valeurs
    assert 'AUDIT création' not in valeurs

    response = admin.get('/audit/export/excel?action=LOGIN')
    feuille = openpyxl.load_workbook(io.BytesIO(response.data)).active
    valeurs = ' '.join(str(c.value or '') for ligne in feuille.iter_rows()
                       for c in ligne)
    assert 'AUDIT login admin' in valeurs
    assert 'AUDIT création' not in valeurs


def test_bouton_export_conserve_les_filtres_affiches(app):
    """Le lien du bouton porte la sélection courante, sans la page."""
    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    page = admin.get('/audit?username=rh&action=LOGIN').data.decode('utf-8')
    assert 'username=rh' in page
    assert 'action=LOGIN' in page


def test_export_reserve_a_ladministrateur(app):
    for username, mot_de_passe in (('rh', 'rh123'),
                                   ('manager', 'manager123'),
                                   ('employe', 'user123')):
        client = app.test_client()
        client.post('/login', data={'username': username, 'password': mot_de_passe})
        assert client.get('/audit/export/excel',
                          follow_redirects=False).status_code in (301, 302), username


def test_utilisateurs_et_journal_partagent_un_deroulant(app):
    """Les deux entrées sont regroupées sous un même menu."""
    import re

    admin = app.test_client()
    admin.post('/login', data={'username': 'admin', 'password': 'admin123'})
    page = admin.get('/').data.decode('utf-8')

    blocs = [m.group(1) for m in
             re.finditer(r'<div class="nav-dropdown">(.*?)</div>\s*</div>', page, re.S)]
    groupe = [b for b in blocs if 'Comptes utilisateurs' in b]
    assert groupe, "aucun déroulant Utilisateurs trouvé"
    assert 'Journal d' in groupe[0]
    assert 'nav-dropdown-toggle' in groupe[0]

    # Le RH garde le déroulant, mais sans le journal : il contourne l'ACL,
    # le masquage automatique ne peut donc pas le masquer à sa place.
    rh = app.test_client()
    rh.post('/login', data={'username': 'rh', 'password': 'rh123'})
    page_rh = rh.get('/').data.decode('utf-8')
    assert 'Comptes utilisateurs' in page_rh
    assert not any('Journal d' in m.group(1) for m in
                   re.finditer(r'<div class="nav-dropdown">(.*?)</div>\s*</div>',
                               page_rh, re.S))
